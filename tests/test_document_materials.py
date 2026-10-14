"""Tests for Word / plain-text materials: extraction, preview, notes.

``.docx`` and ``.txt`` extract text directly; a legacy ``.doc``
is converted to ``.docx`` with a local LibreOffice first. The
inline preview reuses that text (or a LibreOffice-rendered PDF
when LibreOffice is present), and notes generation works from
the extracted text exactly as it does for PDFs and slides.
"""

import io
import uuid

import pytest

from app.services import extractors, office_conversion, preview_service


def _signup(client, email: str) -> dict[str, str]:
    response = client.post(
        "/api/auth/signup",
        json={"email": email, "password": "password123", "full_name": "Tester"},
    )
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _create_course(client, headers: dict[str, str], name: str = "Biology") -> dict:
    response = client.post("/api/courses", json={"name": name}, headers=headers)
    assert response.status_code == 201, response.text
    return response.json()


def _user_id(client, headers: dict[str, str]) -> str:
    response = client.get("/api/auth/me", headers=headers)
    assert response.status_code == 200, response.text
    return response.json()["user"]["id"]


def _upload(client, headers, course_id: str, filename: str, data: bytes, mime: str) -> dict:
    response = client.post(
        f"/api/courses/{course_id}/materials",
        files=[("files", (filename, data, mime))],
        headers=headers,
    )
    assert response.status_code == 201, response.text
    return response.json()[0]


def _docx_bytes(paragraphs: list[str], tables: list[list[list[str]]] | None = None) -> bytes:
    """A real multi-paragraph .docx, with optional tables."""
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for rows in tables or []:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for col_index, cell in enumerate(row):
                table.cell(row_index, col_index).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _seed_material(
    fake,
    user_id: str,
    course_id: str,
    name: str,
    data: bytes,
    *,
    material_type: str = "doc",
    mime: str = "application/octet-stream",
) -> str:
    material_id = str(uuid.uuid4())
    storage_path = f"{user_id}/{course_id}/{material_id}/{name}"
    fake.stored_objects[f"materials/{storage_path}"] = data
    fake.tables.setdefault("materials", []).append(
        {
            "id": material_id,
            "user_id": user_id,
            "course_id": course_id,
            "name": name,
            "type": material_type,
            "kind": "Upload",
            "size_bytes": len(data),
            "mime_type": mime,
            "storage_path": storage_path,
            "status": "ready",
        }
    )
    return material_id


@pytest.fixture(autouse=True)
def _clean_preview_stages():
    preview_service._PREVIEW_STAGES.clear()
    yield
    preview_service._PREVIEW_STAGES.clear()


# ----------------------------------------------------------------------
# .docx extraction
# ----------------------------------------------------------------------

def test_docx_extracts_paragraphs_and_tables(client) -> None:
    headers = _signup(client, "docx@example.com")
    course = _create_course(client, headers)
    content = _docx_bytes(
        ["First paragraph.", "Second paragraph."],
        tables=[[["H1", "H2"], ["a", "b"]]],
    )
    material = _upload(
        client,
        headers,
        course["id"],
        "notes.docx",
        content,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert material["type"] == "doc"

    # The upload route runs extraction inline, so it is already done.
    status = client.get(
        f"/api/materials/{material['id']}/extraction-status", headers=headers
    ).json()
    assert status["status"] == "ready"
    assert status["unit_count"] >= 1

    units = client.get(
        f"/api/materials/{material['id']}/extracted-units", headers=headers
    ).json()
    text = "\n".join(unit["content"] or "" for unit in units)
    assert "First paragraph." in text
    assert "Second paragraph." in text
    assert "H1 | H2" in text  # table text is preserved


def test_docx_preview_shows_extracted_text_without_libreoffice(client) -> None:
    headers = _signup(client, "docxpreview@example.com")
    course = _create_course(client, headers)
    content = _docx_bytes(["The only paragraph."])
    material = _upload(
        client,
        headers,
        course["id"],
        "notes.docx",
        content,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    preview = client.get(
        f"/api/materials/{material['id']}/preview", headers=headers
    ).json()
    # No LibreOffice here, so the extracted text is the inline preview.
    assert preview["mode"] == "text"
    assert "The only paragraph." in (preview["text"] or "")


def test_docx_generates_notes(client) -> None:
    """Notes generation self-heals: extract -> index -> generate."""
    headers = _signup(client, "docxnotes@example.com")
    course = _create_course(client, headers)
    content = _docx_bytes(
        ["Photosynthesis converts light energy into chemical energy."]
    )
    material = _upload(
        client,
        headers,
        course["id"],
        "notes.docx",
        content,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    response = client.post(
        f"/api/courses/{course['id']}/notes/generate",
        json={"material_id": material["id"]},
        headers=headers,
    )
    assert response.status_code == 201, response.text


# ----------------------------------------------------------------------
# .txt extraction and preview
# ----------------------------------------------------------------------

def test_txt_extracts_and_previews_directly(client) -> None:
    headers = _signup(client, "txt@example.com")
    course = _create_course(client, headers)
    content = "Plain text line one.\nPlain text line two.".encode("utf-8")
    material = _upload(
        client, headers, course["id"], "notes.txt", content, "text/plain"
    )

    status = client.get(
        f"/api/materials/{material['id']}/extraction-status", headers=headers
    ).json()
    assert status["status"] == "ready"
    assert status["unit_count"] >= 1

    # Plain text always previews directly — no conversion, no LibreOffice.
    preview = client.get(
        f"/api/materials/{material['id']}/preview", headers=headers
    ).json()
    assert preview["status"] == "ready"
    assert preview["mode"] == "text"
    assert "Plain text line one." in (preview["text"] or "")
    assert "Plain text line two." in (preview["text"] or "")


# ----------------------------------------------------------------------
# Legacy .doc
# ----------------------------------------------------------------------

def test_doc_without_libreoffice_is_rejected_clearly(client) -> None:
    headers = _signup(client, "doc@example.com")
    course = _create_course(client, headers)
    material = _upload(
        client, headers, course["id"], "legacy.doc", b"legacy binary", "application/msword"
    )

    status = client.get(
        f"/api/materials/{material['id']}/extraction-status", headers=headers
    ).json()
    assert status["status"] == "attention"
    assert status["unit_count"] == 0
    assert any("libreoffice" in warning.lower() for warning in status["warnings"])


def test_doc_extracts_via_libreoffice_roundtrip(
    client, fake_supabase, monkeypatch
) -> None:
    """A legacy .doc is converted to .docx locally, then extracted."""
    headers = _signup(client, "docroundtrip@example.com")
    course = _create_course(client, headers)
    user_id = _user_id(client, headers)

    converted = _docx_bytes(["Converted legacy paragraph."])
    monkeypatch.setattr(extractors, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(
        extractors,
        "convert_to",
        lambda soffice, content, filename, target: converted,
    )

    material_id = _seed_material(
        fake_supabase, user_id, course["id"], "legacy.doc", b"legacy binary"
    )
    response = client.post(
        f"/api/materials/{material_id}/extraction", headers=headers
    )
    assert response.status_code == 202, response.text

    units = client.get(
        f"/api/materials/{material_id}/extracted-units", headers=headers
    ).json()
    assert any("Converted legacy paragraph." in (unit["content"] or "") for unit in units)


# ----------------------------------------------------------------------
# .docx preview as a rendered PDF when LibreOffice is present
# ----------------------------------------------------------------------

def test_docx_preview_converts_to_pdf_when_libreoffice_present(
    client, fake_supabase, monkeypatch
) -> None:
    headers = _signup(client, "docxpdf@example.com")
    course = _create_course(client, headers)
    user_id = _user_id(client, headers)
    material_id = _seed_material(
        fake_supabase,
        user_id,
        course["id"],
        "notes.docx",
        _docx_bytes(["Rendered paragraph."]),
    )

    conversions = {"count": 0}

    def fake_convert(soffice, content, filename):
        conversions["count"] += 1
        return b"%PDF-1.4 converted-document"

    monkeypatch.setattr(preview_service, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(preview_service, "convert_to_pdf", fake_convert)

    started = client.post(
        f"/api/materials/{material_id}/preview", headers=headers
    )
    assert started.status_code == 202, started.text
    assert conversions["count"] == 1

    polled = client.get(
        f"/api/materials/{material_id}/preview", headers=headers
    ).json()
    assert polled["status"] == "ready"
    assert polled["mode"] == "pdf"
    assert polled["url"] and "preview.pdf" in polled["url"]

    # Reopening the already-converted material never converts it again.
    again = client.post(
        f"/api/materials/{material_id}/preview", headers=headers
    ).json()
    assert again["status"] == "ready"
    assert conversions["count"] == 1


# ----------------------------------------------------------------------
# Genuine conversion (skipped when LibreOffice is absent)
# ----------------------------------------------------------------------

@pytest.mark.skipif(
    not office_conversion.conversion_available(),
    reason="LibreOffice is not installed in this environment",
)
def test_real_libreoffice_converts_docx_to_pdf() -> None:
    """End to end: a .docx becomes a real PDF."""
    import fitz

    content = _docx_bytes(["Alpha paragraph.", "Beta paragraph."])
    soffice = office_conversion.find_soffice()
    assert soffice
    pdf = office_conversion.convert_to_pdf(soffice, content, "notes.docx")
    assert pdf.startswith(b"%PDF")
    document = fitz.open(stream=pdf, filetype="pdf")
    try:
        assert document.page_count >= 1
        text = document[0].get_text()
        assert "Alpha paragraph." in text
        assert "Beta paragraph." in text
    finally:
        document.close()
