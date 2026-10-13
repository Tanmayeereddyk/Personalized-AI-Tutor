"""Extraction engines for course materials.

Each engine takes the material row + the raw file bytes, extracts text,
tables, transcript segments and/or visual observations, saves extracted
images/frames into the private ``materials`` storage bucket, and returns a
list of unit dicts ready to be persisted into ``extracted_units``.

Provenance is preserved on every unit:

- ``extraction_method`` distinguishes how the content was produced
  (``pdf_text``, ``pdf_table``, ``ocr``, ``pptx_text``, ``video_frame``,
  ``video_transcript``, ...).
- ``metadata["source"]`` distinguishes original text, OCR output, transcripts
  and generated descriptions.
- ``page_number`` / ``slide_number`` / ``start_timestamp_seconds`` /
  ``end_timestamp_seconds`` keep the exact source location.
- ``asset_path`` keeps the storage reference of any extracted image/frame.

Engines that need optional dependencies (python-pptx, whisper, OCR) degrade
gracefully: they raise ``UnsupportedFormatError`` or record a warning instead
of pretending extraction succeeded.
"""

import io
import os
import shutil
import subprocess
import tempfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from fastapi import HTTPException
from supabase import Client

from app.services import extracted_unit_service

BUCKET = "materials"

# Whisper runs fully locally (never an external transcription API). The model
# name and the provenance constants for transcript units live here / in
# extracted_unit_service so the engine, the API and the frontend agree.
WHISPER_MODEL_NAME = "base"

# Cached Whisper model (loaded once per process; see _load_whisper_model).
_WHISPER_MODEL = None


def _find_tesseract() -> str | None:
    """Locate the Tesseract binary so OCR works without a PATH change."""
    found = shutil.which("tesseract")
    if found:
        return found
    if os.name == "nt":
        local_app_data = os.environ.get("LOCALAPPDATA", "")
        candidates = [
            Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
            / "Tesseract-OCR"
            / "tesseract.exe",
            Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"))
            / "Tesseract-OCR"
            / "tesseract.exe",
            Path(local_app_data) / "Programs" / "Tesseract-OCR" / "tesseract.exe",
        ]
        for candidate in candidates:
            if candidate.is_file():
                return str(candidate)
    return None


def _winget_bin(name: str) -> str | None:
    """Find a winget-installed binary without requiring a PATH change."""
    local_app_data = os.environ.get("LOCALAPPDATA", "")
    if not local_app_data:
        return None
    packages = Path(local_app_data) / "Microsoft" / "WinGet" / "Packages"
    if not packages.is_dir():
        return None
    for exe in packages.glob(f"**/bin/{name}"):
        return str(exe)
    return None


def _find_ffmpeg() -> str | None:
    """Locate the ffmpeg binary (PATH first, then a winget install)."""
    return shutil.which("ffmpeg") or _winget_bin("ffmpeg.exe")


def _ensure_ffmpeg_on_path() -> None:
    """Make a winget-installed ffmpeg findable by Whisper (Windows)."""
    if shutil.which("ffmpeg"):
        return
    exe = _winget_bin("ffmpeg.exe")
    if exe:
        os.environ["PATH"] = f"{Path(exe).parent}{os.pathsep}{os.environ.get('PATH', '')}"


def _find_ffprobe() -> str | None:
    """Locate ffprobe (normally shipped beside ffmpeg)."""
    found = shutil.which("ffprobe")
    if found:
        return found
    ffmpeg = _find_ffmpeg()
    if ffmpeg:
        candidate = Path(ffmpeg).with_name("ffprobe.exe" if os.name == "nt" else "ffprobe")
        if candidate.is_file():
            return str(candidate)
    return _winget_bin("ffprobe.exe")


def _has_audio_stream(video_path: str) -> bool | None:
    """True/False when ffprobe can tell, ``None`` when it cannot.

    ``False`` means the file genuinely has no audio stream, so the user can be
    told there is nothing to transcribe instead of showing an empty transcript.
    """
    ffprobe = _find_ffprobe()
    if not ffprobe:
        return None
    try:
        completed = subprocess.run(
            [
                ffprobe,
                "-v",
                "error",
                "-select_streams",
                "a",
                "-show_entries",
                "stream=index",
                "-of",
                "csv=p=0",
                video_path,
            ],
            capture_output=True,
            text=True,
            timeout=60,
        )
    except Exception:  # noqa: BLE001 - probing is best-effort
        return None
    if completed.returncode != 0:
        return None
    return bool(completed.stdout.strip())


class UnsupportedFormatError(Exception):
    """The material cannot be extracted in this environment."""


@dataclass
class ExtractorResult:
    """What one extraction run produces, before persistence."""

    units: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _storage_path(material: dict, relative_path: str) -> str:
    """Storage path for an extracted asset, under the material's own folder."""
    return (
        f"{material['user_id']}/{material['course_id']}/"
        f"{material['id']}/extracted/{relative_path}"
    )


def _store_asset(
    client: Client, material: dict, relative_path: str, data: bytes, content_type: str
) -> str:
    """Upload an extracted image/frame; return its storage path."""
    path = _storage_path(material, relative_path)
    client.storage.from_(BUCKET).upload(
        path,
        data,
        {"content-type": content_type, "upsert": "true"},
    )
    return path


def _base_unit(material: dict, **extra: object) -> dict:
    return {"material_id": material["id"], **extra}


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------

def _render_table(extracted: list[list]) -> str:
    """Render a PDF table as simple text so it survives round-tripping."""
    lines = []
    for row in extracted:
        cells = [str(c).strip() for c in row]
        lines.append(" | ".join(cells))
    return "\n".join(lines)


def _ocr_page(page) -> str | None:
    """OCR one PDF page, or None if OCR is not configured / produced nothing."""
    try:
        import pytesseract  # noqa: PLC0415 - optional dependency
        from PIL import Image  # noqa: PLC0415 - optional dependency
    except ImportError:
        return None
    tesseract = _find_tesseract()
    if tesseract:
        pytesseract.pytesseract.tesseract_cmd = tesseract
    pix = page.get_pixmap(dpi=200)
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    try:
        text = pytesseract.image_to_string(image) or ""
    except Exception:  # noqa: BLE001 - OCR tools are external
        return None
    text = text.strip()
    return text or None


def extract_pdf(client: Client, material: dict, content: bytes) -> ExtractorResult:
    """Extract per-page text, tables and a page image for each PDF page."""
    try:
        import pymupdf  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - guarded import
        raise UnsupportedFormatError(
            "PDF extraction requires PyMuPDF, which is not installed."
        ) from exc

    try:
        document = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:  # noqa: BLE001 - malformed PDFs are a user error
        raise UnsupportedFormatError(f"Could not open the PDF: {exc}") from exc

    units: list[dict] = []
    warnings: list[str] = []

    try:
        for page_number, page in enumerate(document, start=1):
            text = (page.get_text("text") or "").strip()

            table_texts: list[str] = []
            try:
                for table in page.find_tables().tables:
                    rendered = _render_table(table.extract())
                    if rendered.strip():
                        table_texts.append(rendered)
            except Exception:  # noqa: BLE001 - table detection is best-effort
                pass

            method = "pdf_text"
            source = "text_layer"
            content_text = "\n\n".join(
                part for part in [text, *table_texts] if part.strip()
            )

            if not content_text:
                # No text layer on this page. OCR is the fallback when
                # configured; otherwise we keep the visual observation and
                # record the gap as a warning instead of fabricating content.
                ocr = _ocr_page(page)
                if ocr:
                    content_text = ocr
                    method = "ocr"
                    source = "ocr"
                else:
                    warnings.append(
                        f"PDF page {page_number} has no text layer and OCR is "
                        "not configured; only its image was kept."
                    )

            asset_path: str | None = None
            try:
                pix = page.get_pixmap(dpi=110)
                asset_path = _store_asset(
                    client,
                    material,
                    f"pdf/page-{page_number}.png",
                    pix.tobytes("png"),
                    "image/png",
                )
            except Exception as exc:  # noqa: BLE001 - storage is external
                warnings.append(
                    f"PDF page {page_number}: could not save the page image ({exc})"
                )

            units.append(
                _base_unit(
                    material,
                    content_type="text",
                    content=content_text or None,
                    page_number=page_number,
                    asset_path=asset_path,
                    extraction_method=method,
                    status="extracted",
                    unit_order=page_number * 10,
                    metadata={"source": source},
                )
            )

            for index, table_text in enumerate(table_texts, start=1):
                units.append(
                    _base_unit(
                        material,
                        content_type="table",
                        content=table_text,
                        page_number=page_number,
                        extraction_method="pdf_table",
                        status="extracted",
                        unit_order=page_number * 10 + index,
                        metadata={"source": "tables"},
                    )
                )
    finally:
        document.close()

    return ExtractorResult(units=units, warnings=warnings)


# ----------------------------------------------------------------------
# PPT / PPTX
# ----------------------------------------------------------------------

def _collect_slide_text(slide) -> list[str]:
    parts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = (shape.text_frame.text or "").strip()
            if text:
                parts.append(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return parts


def extract_pptx(client: Client, material: dict, content: bytes) -> ExtractorResult:
    """Extract per-slide text from PPTX files (legacy .ppt is rejected)."""
    try:
        from pptx import Presentation  # type: ignore[import-not-found]
    except ImportError as exc:
        raise UnsupportedFormatError(
            "PPT/PPTX extraction requires the 'python-pptx' package, "
            "which is not installed.",
        ) from exc

    name = (material.get("name") or "").lower()
    if name.endswith(".ppt") and not name.endswith(".pptx"):
        raise UnsupportedFormatError(
            "Legacy .ppt files are not supported; convert the file to .pptx first.",
        )

    try:
        presentation = Presentation(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 - malformed files are a user error
        raise UnsupportedFormatError(f"Could not open the presentation: {exc}") from exc

    units: list[dict] = []
    warnings: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        parts = _collect_slide_text(slide)
        content_text = "\n\n".join(parts)

        # Slide rendering requires a desktop office suite (LibreOffice), which
        # is not assumed on this server; text extraction is still complete.
        units.append(
            _base_unit(
                material,
                content_type="text",
                content=content_text or None,
                slide_number=index,
                extraction_method="pptx_text",
                status="extracted",
                unit_order=index * 10,
                metadata={"source": "pptx_text"},
            )
        )
        if not content_text:
            warnings.append(
                f"Slide {index} has no extractable text (it may be an image)."
            )

    return ExtractorResult(units=units, warnings=warnings)


# ----------------------------------------------------------------------
# Video
# ----------------------------------------------------------------------

def _transcribe_available() -> bool:
    """True when genuine transcription is possible in this environment."""
    if not _find_ffmpeg():
        return False
    try:
        # The PyPI package is "openai-whisper"; it imports as "whisper".
        import whisper  # noqa: PLC0415, F401 - optional dependency
    except ImportError:
        return False
    return True


def _remove_file(path: str | None) -> None:
    """Best-effort temp-file cleanup (never raises)."""
    if not path:
        return
    try:
        os.unlink(path)
    except OSError:  # noqa: S110 - temp cleanup is best-effort
        pass


def _load_whisper_model():
    """Load (and cache) the local Whisper model.

    Kept as a process-wide singleton so a long lecture does not pay the
    model-load cost again on every extraction. Tests monkeypatch this.
    """
    global _WHISPER_MODEL
    if _WHISPER_MODEL is None:
        import whisper  # noqa: PLC0415 - optional dependency

        _WHISPER_MODEL = whisper.load_model(WHISPER_MODEL_NAME)
    return _WHISPER_MODEL


def _extract_audio_track(video_path: str) -> tuple[str | None, str | None]:
    """Extract a video's audio to a temporary 16 kHz mono WAV with ffmpeg.

    Returns ``(audio_path, warning)``. On success ``warning`` is ``None``; on
    failure ``audio_path`` is ``None`` and ``warning`` explains why, so the
    user is told instead of being shown an empty transcript. The caller owns
    the returned file and must delete it.
    """
    ffmpeg = _find_ffmpeg()
    if not ffmpeg:
        return None, (
            "Transcript extraction requires the ffmpeg binary, which was not "
            "found; only video frames were extracted."
        )

    handle, audio_path = tempfile.mkstemp(suffix=".wav")
    os.close(handle)
    command = [
        ffmpeg,
        "-nostdin",
        "-y",
        "-i",
        video_path,
        "-vn",  # drop the video stream
        "-ac",
        "1",  # mono
        "-ar",
        "16000",  # 16 kHz — what Whisper expects
        "-f",
        "wav",
        audio_path,
    ]
    try:
        completed = subprocess.run(  # noqa: S603 - fixed argv, no shell
            command, capture_output=True, text=True, timeout=1800
        )
    except Exception as exc:  # noqa: BLE001 - external process
        _remove_file(audio_path)
        return None, f"The audio track could not be extracted from this video ({exc})."

    if completed.returncode != 0 or not os.path.exists(audio_path):
        _remove_file(audio_path)
        return None, (
            "The audio track could not be extracted from this video; it may "
            "have no audio or use a codec ffmpeg cannot read."
        )
    # A WAV with only its 44-byte header means ffmpeg found no samples.
    if os.path.getsize(audio_path) <= 44:
        _remove_file(audio_path)
        return None, "This video has no audio track, so there is no speech to transcribe."
    return audio_path, None


def _transcribe_audio(audio_path: str) -> dict:
    """Transcribe an extracted audio file locally with Whisper.

    No audio ever leaves the machine: this runs the Whisper model in-process.
    """
    _ensure_ffmpeg_on_path()
    model = _load_whisper_model()
    return model.transcribe(audio_path)


def _transcript_units(
    material: dict, transcription: dict, first_order: int
) -> list[dict]:
    """Map Whisper segments to unit rows, keeping timestamps and language."""
    language = transcription.get("language")
    units: list[dict] = []
    for index, segment in enumerate(transcription.get("segments") or [], start=1):
        text = (segment.get("text") or "").strip()
        if not text:
            continue
        start = float(segment.get("start") or 0.0)
        end = float(segment.get("end") or 0.0)
        units.append(
            _base_unit(
                material,
                content_type="text",
                content=text,
                start_timestamp_seconds=round(start, 2),
                end_timestamp_seconds=round(end, 2),
                extraction_method=extracted_unit_service.TRANSCRIPT_METHOD,
                status="extracted",
                unit_order=first_order + index,
                metadata={
                    "source": extracted_unit_service.AUDIO_TRANSCRIPT_SOURCE,
                    "kind": extracted_unit_service.TRANSCRIPT_KIND,
                    "language": language,
                    "segment": index,
                },
            )
        )
    return units


def extract_transcript_units(
    material: dict,
    video_path: str,
    *,
    on_stage: Callable[[str], None] | None = None,
) -> tuple[list[dict], list[str]]:
    """Extract and locally transcribe a video's speech.

    Returns ``(units, warnings)``. Nothing is ever sent to an external API:
    the audio is pulled out to a temporary file with ffmpeg, transcribed with
    the local Whisper model and the temp file is always removed. Every failure
    (no audio stream, silent or unclear speech, unsupported codec, a crashing
    transcription) becomes a warning instead of an empty transcript.
    """
    units: list[dict] = []
    warnings: list[str] = []

    if not _transcribe_available():
        warnings.append(
            "Transcript extraction requires 'openai-whisper' and ffmpeg, "
            "which are not installed; only video frames were extracted."
        )
        return units, warnings

    if _has_audio_stream(video_path) is False:
        # ffprobe is authoritative: there is genuinely nothing to hear.
        warnings.append(
            "This video has no audio track, so there is no speech to transcribe."
        )
        return units, warnings

    if on_stage:
        on_stage("extracting_audio")
    audio_path, audio_warning = _extract_audio_track(video_path)
    if audio_path is None:
        warnings.append(
            audio_warning or "The audio track could not be extracted from this video."
        )
        return units, warnings

    try:
        if on_stage:
            on_stage("transcribing")
        try:
            transcription = _transcribe_audio(audio_path)
        except Exception as exc:  # noqa: BLE001 - best-effort, warn only
            warnings.append(f"Speech transcription failed ({exc}).")
        else:
            units = _transcript_units(material, transcription, 10_000_000)
            if not units:
                warnings.append(
                    "The audio was transcribed but no clear speech was "
                    "recognized, so there is no transcript."
                )
    finally:
        _remove_file(audio_path)  # never leave the extracted audio behind

    return units, warnings


def extract_video(
    client: Client,
    material: dict,
    content: bytes,
    frame_interval_seconds: float = 30.0,
    max_frames: int = 120,
    *,
    on_stage: Callable[[str], None] | None = None,
) -> ExtractorResult:
    """Extract keyframes and a timestamped speech transcript from a video.

    Frames are decoded with OpenCV; the audio track is extracted with ffmpeg
    and transcribed locally with Whisper (never an external API). Every
    failure is recorded as a warning instead of yielding an empty transcript,
    so the UI can tell the user exactly what happened.
    """
    try:
        import cv2  # type: ignore[import-not-found]
    except ImportError as exc:
        raise UnsupportedFormatError(
            "Video extraction requires OpenCV (opencv-python), which is not installed.",
        ) from exc

    tmp_path: str | None = None
    units: list[dict] = []
    warnings: list[str] = []

    try:
        with tempfile.NamedTemporaryFile(suffix=".mp4", delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name

        capture = cv2.VideoCapture(tmp_path)
        if not capture.isOpened():
            raise UnsupportedFormatError("Could not decode the video file.")

        try:
            fps = float(capture.get(cv2.CAP_PROP_FPS)) or 30.0
            total_frames = int(capture.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
            duration = total_frames / fps if total_frames > 0 and fps > 0 else None

            if duration is None or duration <= 0:
                warnings.append(
                    "Could not determine the video duration; no frames were extracted."
                )
            else:
                # Sample evenly across the whole video, capping the total
                # number of frames so a two-hour lecture doesn't flood storage.
                step = max(frame_interval_seconds, duration / max_frames) if max_frames > 0 else frame_interval_seconds
                timestamp = 0.0
                while timestamp < duration:
                    frame_number = int(round(timestamp * fps))
                    capture.set(cv2.CAP_PROP_POS_FRAMES, frame_number)
                    ok, frame = capture.read()
                    if not ok:
                        warnings.append(
                            f"Skipped the frame at {timestamp:.0f}s (could not be decoded)."
                        )
                    else:
                        ok_encoded, buffer = cv2.imencode(
                            ".jpg", frame, [int(cv2.IMWRITE_JPEG_QUALITY), 80]
                        )
                        if not ok_encoded:
                            warnings.append(
                                f"Skipped the frame at {timestamp:.0f}s (could not be compressed)."
                            )
                        else:
                            asset_path = _store_asset(
                                client,
                                material,
                                f"video/frame-{int(timestamp)}s.jpg",
                                buffer.tobytes(),
                                "image/jpeg",
                            )
                            units.append(
                                _base_unit(
                                    material,
                                    content_type="image",
                                    content=None,
                                    start_timestamp_seconds=round(timestamp, 2),
                                    end_timestamp_seconds=round(
                                        min(timestamp + step, duration), 2
                                    ),
                                    asset_path=asset_path,
                                    extraction_method="video_frame",
                                    status="extracted",
                                    unit_order=1 + int(round(timestamp * 1000)),
                                    metadata={"source": "frame"},
                                )
                            )
                    timestamp += step
        finally:
            capture.release()

        transcript_units, transcript_warnings = extract_transcript_units(
            material, tmp_path, on_stage=on_stage
        )
        units.extend(transcript_units)
        warnings.extend(transcript_warnings)
    finally:
        _remove_file(tmp_path)

    return ExtractorResult(units=units, warnings=warnings)


# ----------------------------------------------------------------------
# Registry
# ----------------------------------------------------------------------

def run_extractor(
    client: Client,
    material: dict,
    content: bytes,
    *,
    on_stage: Callable[[str], None] | None = None,
) -> ExtractorResult:
    """Dispatch to the right engine for a material row.

    ``on_stage`` is an optional progress callback used by the video engine to
    report which phase it is in (``extracting_audio`` / ``transcribing``). It
    is ignored by the engines that have no long-running stages.
    """
    material_type = (material.get("type") or "doc").lower()
    if material_type == "pdf":
        return extract_pdf(client, material, content)
    if material_type == "ppt":
        return extract_pptx(client, material, content)
    if material_type == "video":
        return extract_video(client, material, content, on_stage=on_stage)
    return ExtractorResult(
        units=[],
        warnings=[
            f"Extraction is not implemented for this file type ('{material_type}')."
        ],
    )
    return ExtractorResult(
        units=[],
        warnings=[
            f"Extraction is not implemented for this file type ('{material_type}')."
        ],
    )


def download_material(client: Client, material: dict) -> bytes:
    """Read the material's stored file bytes (server-side only)."""
    path = material.get("storage_path")
    if not path:
        raise HTTPException(status_code=404, detail="The source file is unavailable")
    try:
        return client.storage.from_(BUCKET).download(path)
    except Exception as exc:  # noqa: BLE001 - storage failures are external
        raise HTTPException(
            status_code=502,
            detail="Could not read the source file from storage",
        ) from exc


def summary_unit(material: dict, warnings: list[str]) -> dict | None:
    """A single unit that records extraction warnings (only when present).

    Never inserted to fake success — it only exists when a run really had
    partial failures or missing capabilities.
    """
    if not warnings:
        return None
    return _base_unit(
        material,
        content_type=extracted_unit_service.SUMMARY_CONTENT_TYPE,
        content=None,
        extraction_method=f"{material.get('type', 'doc')}:summary",
        status="warning",
        unit_order=0,
        metadata={
            "kind": extracted_unit_service.SUMMARY_KIND,
            "warnings": warnings,
        },
    )